"""
Amazon segment variance report: one-step refresh.

Adding a quarter:
    1. Append three rows to data/segments.csv (one per segment) and one row to
       data/consolidated_check.csv, straight from the new 10-Q / earnings release.
    2. Run:  python refresh.py

The script rebuilds amazon.db from scratch, stops if the segments do not tie to
the reported consolidated totals, then writes output/segment_report.xlsx and
output/segment_chart.png.
"""
import csv
import sqlite3
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
DB = ROOT / "amazon.db"
OUT = ROOT / "output"
SEGMENTS = ["North America", "International", "AWS"]
COLORS = {"North America": "#2a78d6", "International": "#eb6834", "AWS": "#1baf7a"}


def load_csv(con, table, path):
    with open(path, newline="") as f:
        rows = list(csv.DictReader(f))
    cols = list(rows[0].keys())
    con.executemany(
        f"INSERT INTO {table} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
        [tuple(r[c] for c in cols) for r in rows],
    )


def build_db():
    DB.unlink(missing_ok=True)
    con = sqlite3.connect(DB)
    for sql in sorted((ROOT / "sql").glob("*.sql")):
        con.executescript(sql.read_text())
    load_csv(con, "segment_results", ROOT / "data" / "segments.csv")
    load_csv(con, "consolidated_check", ROOT / "data" / "consolidated_check.csv")
    con.commit()
    return con


def check_tieout(con):
    tie = pd.read_sql("SELECT * FROM v_tieout", con)
    bad = tie[(tie.sales_diff != 0) | (tie.oi_diff != 0)]
    missing = set(pd.read_sql("SELECT DISTINCT fiscal_quarter FROM segment_results", con).fiscal_quarter) - set(tie.fiscal_quarter)
    if len(bad) or missing:
        print("STOP: segment rows do not tie to reported totals.")
        print(bad.to_string(index=False))
        if missing:
            print("No consolidated_check row for:", sorted(missing))
        sys.exit(1)
    return tie


def check_second_source():
    """Compare every figure to the same quarter as printed in a different filing."""
    d = pd.read_csv(ROOT / "data" / "segments.csv")
    c = pd.read_csv(ROOT / "data" / "crosscheck_sources.csv")
    m = d.merge(c, on=["fiscal_quarter", "segment"], how="left", suffixes=("", "_chk"))
    unchecked = m[m.check_source.isna()]
    bad = m[m.check_source.notna() & ((m.net_sales_musd != m.net_sales_musd_chk)
                                      | (m.operating_income_musd != m.operating_income_musd_chk))]
    if len(bad):
        print("STOP: figures differ from the second-source filing.")
        print(bad[["fiscal_quarter", "segment", "net_sales_musd", "net_sales_musd_chk",
                   "operating_income_musd", "operating_income_musd_chk", "check_source"]].to_string(index=False))
        sys.exit(1)
    return len(m) - len(unchecked), unchecked


def make_chart(df):
    oi = df.pivot(index="fiscal_quarter", columns="segment", values="oi")[SEGMENTS] / 1000
    margin = df.pivot(index="fiscal_quarter", columns="segment", values="op_margin_pct")[SEGMENTS]
    labels = [f"Q{q[-1]}\n'{q[2:4]}" for q in oi.index]
    x = range(len(oi))

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(13, 5.2), facecolor="#fcfcfb")
    w = 0.26
    for i, seg in enumerate(SEGMENTS):
        a1.bar([p + (i - 1) * w for p in x], oi[seg], width=w - 0.03,
               color=COLORS[seg], label=seg)
    a1.set_title("Operating income by segment ($B)", loc="left", fontweight="bold", color="#0b0b0b")
    a1.axhline(0, color="#52514e", linewidth=0.8)

    for seg in SEGMENTS:
        a2.plot(list(x), margin[seg], color=COLORS[seg], linewidth=2, marker="o", markersize=4, label=seg)
        a2.annotate(f"{seg}  {margin[seg].iloc[-1]:.1f}%", (len(x) - 1, margin[seg].iloc[-1]),
                    xytext=(8, 0), textcoords="offset points", va="center", color="#0b0b0b", fontsize=9)
    a2.set_title("Operating margin by segment (%)", loc="left", fontweight="bold", color="#0b0b0b")
    a2.axhline(0, color="#52514e", linewidth=0.8)
    a2.set_xlim(-0.5, len(x) + 2.2)

    for ax in (a1, a2):
        ax.set_facecolor("#fcfcfb")
        ax.set_xticks(list(x))
        ax.set_xticklabels(labels, color="#52514e", fontsize=8.5)
        ax.tick_params(axis="y", colors="#52514e")
        ax.grid(axis="y", color="#e6e5e0", linewidth=0.8)
        ax.set_axisbelow(True)
        for s in ("top", "right", "left"):
            ax.spines[s].set_visible(False)
        ax.spines["bottom"].set_color("#c3c2b7")
    a1.legend(frameon=False, loc="upper left", fontsize=9)

    first, last = oi.index[0], oi.index[-1]
    fig.suptitle(f"Amazon segment performance, {first[:4]} Q{first[-1]} to {last[:4]} Q{last[-1]}",
                 x=0.01, ha="left", fontsize=14, fontweight="bold", color="#0b0b0b")
    fig.text(0.01, 0.005, "Source: Amazon.com 10-Q, 10-K and earnings release (8-K Ex. 99.1) filings. Segment figures tie to reported consolidated totals.",
             fontsize=8, color="#52514e")
    fig.tight_layout(rect=(0, 0.03, 1, 0.97))
    path = OUT / "segment_chart.png"
    fig.savefig(path, dpi=200, facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def style_workbook(path, wrap_sheets=("Commentary",)):
    wb = load_workbook(path)
    head = PatternFill("solid", fgColor="1F3A5F")
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, (int, float)):
                    hdr = str(ws.cell(1, c.column).value)
                    c.number_format = "0.0" if ("pct" in hdr or "%" in hdr) else "#,##0;(#,##0)"
                if ws.title in wrap_sheets:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    con = build_db()
    tie = check_tieout(con)
    n_checked, unchecked = check_second_source()

    analysis = pd.read_sql("SELECT * FROM v_segment_analysis", con)
    swings = pd.read_sql("SELECT * FROM v_biggest_swings LIMIT 10", con)
    raw = pd.read_sql("SELECT * FROM segment_results ORDER BY fiscal_quarter, segment", con)
    commentary = pd.read_csv(ROOT / "data" / "commentary.csv")

    latest = analysis.fiscal_quarter.max()
    summary = analysis[analysis.fiscal_quarter == latest][
        ["segment", "sales", "oi", "op_margin_pct", "sales_qoq_pct", "sales_yoy_pct",
         "oi_yoy_change", "share_of_sales_pct", "share_of_oi_pct"]
    ].rename(columns={"sales": "net_sales_musd", "oi": "operating_income_musd"})

    xlsx = OUT / "segment_report.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name=f"Summary {latest}", index=False)
        commentary.to_excel(xw, sheet_name="Commentary", index=False)
        swings.to_excel(xw, sheet_name="Biggest swings", index=False)
        analysis.to_excel(xw, sheet_name="Segment analysis", index=False)
        tie.to_excel(xw, sheet_name="Tie-out", index=False)
        raw.to_excel(xw, sheet_name="Raw data", index=False)
    style_workbook(xlsx)

    chart = make_chart(analysis)
    unverified = (raw.verified != "Y").sum()
    print(f"Refreshed through {latest}: {len(raw)} rows, all quarters tie to reported totals.")
    print(f"  Second-source check: {n_checked} of {len(raw)} figures match a separate filing.")
    print(f"  {xlsx.relative_to(ROOT)}\n  {chart.relative_to(ROOT)}")
    if len(unchecked):
        print("  Not yet second-sourced:", ", ".join(sorted(set(unchecked.fiscal_quarter))))
    if unverified:
        print(f"  Note: {unverified} rows still marked verified=N in data/segments.csv")


if __name__ == "__main__":
    main()
